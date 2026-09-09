from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "BRAIN智能工作台V1.0-软著源代码-补文件头.docx"
PAGES = 60
LINES_PER_PAGE = 60
HALF = PAGES // 2
HALF_SOURCE_LINES = 1785
MAX_SOURCE_LINE = 100


def source_files() -> list[Path]:
    listed = subprocess.check_output(
        ["git", "ls-files", "shared", "platforms/windows"],
        cwd=ROOT,
        text=True,
        encoding="utf-8",
        errors="ignore",
    ).splitlines()
    out: list[Path] = []
    for value in listed:
        path = ROOT / value
        normalized = value.replace("\\", "/")
        if path.suffix.lower() not in {".ts", ".tsx", ".js", ".jsx"}:
            continue
        if re.search(r"(^|/)(build|tmp|node_modules|fixtures?|generated)(/|$)", normalized, re.I):
            continue
        if re.search(r"\.(test|spec)\.", path.name, re.I):
            continue
        out.append(path)
    return sorted(out, key=lambda p: str(p).lower())


def clean_source(path: Path) -> list[str]:
    lines: list[str] = []
    in_block = False
    for raw in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        text = raw.rstrip()
        stripped = text.strip()
        if in_block:
            if "*/" in stripped:
                in_block = False
            continue
        if stripped.startswith("/*"):
            if "*/" not in stripped[2:]:
                in_block = True
            continue
        if not stripped or stripped.startswith("//") or stripped.startswith("*"):
            continue
        text = re.sub(r"\s+//.*$", "", text).rstrip()
        if not text:
            continue
        if len(text) > MAX_SOURCE_LINE:
            continue
        if re.search(r"作者|地址|电话|邮箱|copyright|著作权", text, re.I):
            continue
        lines.append(text)
    return lines


def collect_entries() -> list[tuple[Path, str]]:
    stream: list[tuple[Path, str]] = []
    for path in source_files():
        for line in clean_source(path):
            stream.append((path, line))
    required = HALF_SOURCE_LINES * 2
    if len(stream) < required:
        raise RuntimeError(f"source pool has {len(stream)} lines; {required} required")
    return stream[:HALF_SOURCE_LINES] + stream[-HALF_SOURCE_LINES:]


def build() -> None:
    entries = collect_entries()
    if len(entries) != HALF_SOURCE_LINES * 2:
        raise RuntimeError(f"expected {HALF_SOURCE_LINES * 2} source lines, got {len(entries)}")

    file_blocks = 0
    prev_path: Path | None = None
    for path, _ in entries:
        if path != prev_path:
            file_blocks += 1
            prev_path = path

    title_lines = 3
    total_paragraphs = len(entries) + file_blocks + title_lines
    if total_paragraphs != PAGES * LINES_PER_PAGE:
        raise RuntimeError(
            f"expected {PAGES * LINES_PER_PAGE} total paragraphs, got {total_paragraphs} "
            f"(source={len(entries)}, files={file_blocks}, title={title_lines})"
        )

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.79)
    section.right_margin = Inches(0.79)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.header_distance = Inches(0.50)
    section.footer_distance = Inches(0.50)

    header = section.header.paragraphs[0]
    header.text = "BRAIN智能工作台 V1.0"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(9)

    normal = doc.styles["Normal"]
    normal.font.name = "Courier New"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def add_line(text: str, *, bold: bool = False, center: bool = False, spacing_after: int = 0) -> None:
        nonlocal paragraph_count
        if paragraph_count and paragraph_count % LINES_PER_PAGE == 0:
            doc.add_page_break()
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = False
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(8.5)
        run.bold = bold
        if spacing_after:
            paragraph.paragraph_format.space_after = Pt(spacing_after)
        paragraph_count += 1

    paragraph_count = 0
    add_line("==============================================================", bold=True)
    add_line("BRAIN智能工作台 V1.0 前后端源代码", bold=True)
    add_line("==============================================================", bold=True, spacing_after=6)

    prev_path = None
    for path, line in entries:
        if path != prev_path:
            rel = path.relative_to(ROOT).as_posix()
            add_line(f"----- {rel} -----", bold=True, spacing_after=3)
            prev_path = path
        add_line(line)

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.company = ""
    props.comments = ""
    doc.save(OUTPUT)
    print(
        f"created {OUTPUT} with {len(entries)} source lines, "
        f"{file_blocks} file headers across {PAGES} explicit pages"
    )


if __name__ == "__main__":
    build()
